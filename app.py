from flask import Flask, request, render_template, send_file
from pptx import Presentation
from pptx.util import Pt, Emu, Inches
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn, nsmap
from pptx.oxml import parse_xml
import smtplib
from email.mime.multipart import MIMEMultipart
from email.mime.base import MIMEBase
from email.mime.text import MIMEText
from email import encoders
import os, io, tempfile
from lxml import etree
from PIL import Image as PILImage

app = Flask(__name__)
UPLOAD_FOLDER = tempfile.mkdtemp()

# ── 색상 ────────────────────────────────────────────────
BLACK       = RGBColor(0x00, 0x00, 0x00)
WHITE       = RGBColor(0xFF, 0xFF, 0xFF)
HANSOL_BLUE = RGBColor(0x00, 0x35, 0x7A)
GRAY_HDR    = RGBColor(0xD9, 0xD9, 0xD9)   # 표 헤더 / 구분 배경
BORDER_CLR  = RGBColor(0x80, 0x80, 0x80)

I = 914400
def px(inch): return int(inch * I)
def cm(c): return int(c * 360000)   # cm → EMU
def rgb_s(c): return str(c)   # RGBColor → 'RRGGBB'

# ── XML 헬퍼 ────────────────────────────────────────────
def _ns(tag): return f'{{http://schemas.openxmlformats.org/drawingml/2006/main}}{tag}'

def remove_table_style(table):
    """기본 테마 스타일 제거 → 모든 셀 색상을 직접 제어 가능하게"""
    tblPr = table._tbl.find(_ns('tblPr'))
    if tblPr is None:
        tblPr = etree.SubElement(table._tbl, _ns('tblPr'))
    for ts in tblPr.findall(_ns('tableStyleId')):
        tblPr.remove(ts)
    ts = etree.SubElement(tblPr, _ns('tableStyleId'))
    ts.text = '{00000000-0000-0000-0000-000000000000}'

def set_bg(cell, color: RGBColor):
    tcPr = cell._tc.get_or_add_tcPr()
    for sf in tcPr.findall(_ns('solidFill')):
        tcPr.remove(sf)
    sf  = etree.SubElement(tcPr, _ns('solidFill'))
    sc  = etree.SubElement(sf,   _ns('srgbClr'))
    sc.set('val', rgb_s(color))

def set_border(cell, color=BLACK, pt=0.75):
    w = int(pt * 12700)
    tcPr = cell._tc.get_or_add_tcPr()
    for side in ['lnL','lnR','lnT','lnB']:
        tag = _ns(side)
        for old in tcPr.findall(tag):
            tcPr.remove(old)
        ln = etree.SubElement(tcPr, tag)
        ln.set('w', str(w))
        sf = etree.SubElement(ln, _ns('solidFill'))
        sc = etree.SubElement(sf, _ns('srgbClr'))
        sc.set('val', rgb_s(color))

def cell_write(cell, lines, sz=9, bold=False, color=BLACK,
               align=PP_ALIGN.LEFT, font='맑은 고딕'):
    """lines: str 또는 [(text, sz, bold), ...] 또는 [str, ...]"""
    tf = cell.text_frame
    tf.word_wrap = True
    tf.clear()
    # 내부 여백 최소화
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    for tag in ['marL','marR','marT','marB']:
        tcPr.set(tag, str(px(0.04)))

    if isinstance(lines, str):
        lines = [lines]

    for i, line in enumerate(lines):
        if isinstance(line, tuple):
            text, lsz, lbold = line[0], line[1], line[2] if len(line)>2 else bold
        else:
            text, lsz, lbold = line, sz, bold

        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        # 줄 간격
        pPr = p._p.get_or_add_pPr()
        lnSpc = etree.SubElement(pPr, _ns('lnSpc'))
        spcPct = etree.SubElement(lnSpc, _ns('spcPct'))
        spcPct.set('val', '100000')
        spBef = etree.SubElement(pPr, _ns('spcBef'))
        spcPts_el = etree.SubElement(spBef, _ns('spcPts'))
        spcPts_el.set('val', '0')

        r = p.add_run()
        r.text = text
        r.font.size = Pt(lsz)
        r.font.bold = lbold
        r.font.color.rgb = color
        r.font.name = font

def add_rect(slide, l, t, w, h, fill=None, line=None, lpt=0.75):
    s = slide.shapes.add_shape(1, px(l), px(t), px(w), px(h))
    if fill:
        s.fill.solid(); s.fill.fore_color.rgb = fill
    else:
        s.fill.background()
    if line:
        s.line.color.rgb = line; s.line.width = int(lpt*12700)
    else:
        s.line.fill.background()
    return s

def add_tb(slide, l, t, w, h, text='', sz=9, bold=False,
           color=BLACK, align=PP_ALIGN.LEFT, font='맑은 고딕', wrap=True):
    shape = slide.shapes.add_textbox(px(l), px(t), px(w), px(h))
    tf = shape.text_frame
    tf.word_wrap = wrap
    if text:
        p = tf.paragraphs[0]
        p.alignment = align
        r = p.add_run()
        r.text = text; r.font.size = Pt(sz)
        r.font.bold = bold; r.font.color.rgb = color; r.font.name = font
    return tf

def insert_pic(slide, path, l, t, w, h):
    try:
        img = PILImage.open(path)
        iw, ih = img.size
        ratio = ih / iw
        wi, hi = w, w * ratio
        if hi > h:
            hi = h; wi = h / ratio
        ox = (w - wi) / 2; oy = (h - hi) / 2
        slide.shapes.add_picture(path, px(l+ox), px(t+oy), px(wi), px(hi))
    except Exception as e:
        print(f'이미지 오류: {e}')


# ════════════════════════════════════════════════════════
#  PPT 생성 메인
# ════════════════════════════════════════════════════════
def create_report_ppt(data: dict,
                      image_paths: list,
                      cause_image_paths: list = None) -> bytes:
    cause_image_paths = cause_image_paths or []

    prs = Presentation()
    prs.slide_width  = px(13.33)
    prs.slide_height = px(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    SW = 33.87/2.54   # 슬라이드 너비 (inch) — LAYOUT_WIDE 기준
    SH = 19.05/2.54   # 슬라이드 높이 (inch)
    M  = 0.15/2.54    # 기본 여백

    # ────────────────────────────────────────────────────
    # 1. 헤더 영역  (단위: cm → EMU via cm())
    # ────────────────────────────────────────────────────

    # 1-1. 팀/날짜 박스  x=0.15 y=0.12 w=2.82 h=1.15
    s = slide.shapes.add_shape(1, cm(0.15), cm(0.12), cm(2.82), cm(1.15))
    s.fill.solid(); s.fill.fore_color.rgb = WHITE
    s.line.color.rgb = BLACK; s.line.width = int(0.75*12700)
    tf = s.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = cm(0.1)
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = data['team']
    r.font.size = Pt(9); r.font.bold = True
    r.font.color.rgb = BLACK; r.font.name = '맑은 고딕'
    p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
    r2 = p2.add_run(); r2.text = data['date_str']
    r2.font.size = Pt(8); r2.font.color.rgb = BLACK; r2.font.name = '맑은 고딕'

    # 1-2. 제목 박스  x=3.22 y=0.18 w=23.24 h=1.03
    add_tb(slide, 3.22/2.54, 0.18/2.54, 23.24/2.54, 1.03/2.54,
           text=f'■  {data["title"]}',
           sz=15, bold=True, color=BLACK, align=PP_ALIGN.LEFT)

    # 1-3. Hansol 로고 (우측 상단, 제목 오른쪽 남은 공간)
    LOGO_IMG = os.path.join(os.path.dirname(__file__), '한솔로고.jpg')
    if os.path.exists(LOGO_IMG):
        insert_pic(slide, LOGO_IMG, (33.87-6.0)/2.54, 0.12/2.54, 6.0/2.54, 1.15/2.54)

    # 1-4. 헤더 하단 가로 구분선  (제목 박스 바로 아래)
    sep_y = (0.18 + 1.03 + 0.12) / 2.54
    add_rect(slide, 0.15/2.54, sep_y, 33.5/2.54, 0.03/2.54, fill=BLACK)

    # ────────────────────────────────────────────────────
    # 2. ■ 개요
    # ────────────────────────────────────────────────────
    # 2-1. "■ 개요" 라벨  x=0.19 y=1.59 w=6.06 h=0.86
    add_rect(slide, 0.19/2.54, 1.59/2.54, 6.06/2.54, 0.86/2.54,
             fill=WHITE, line=None)
    add_tb(slide, 0.30/2.54, 1.65/2.54, 5.8/2.54, 0.72/2.54,
           text='■ 개요', sz=14, bold=True)

    # 2-2. 일시/지종/중지/내용 칸  x=0.84 y=2.47 w=16.3 h=2.49
    OV_X = 0.84/2.54; OV_Y = 2.47/2.54
    OV_W = 16.3/2.54; OV_H = 2.49/2.54
    # 테두리 박스
    add_rect(slide, OV_X, OV_Y, OV_W, OV_H, fill=WHITE, line=BLACK, lpt=0.75)
    LINE_H = OV_H / 4
    items = [
        f'○ 일    시 : {data["datetime"]}',
        f'○ 지    종 : {data["jijong"]}',
        f'○ 중    지 : {data["jungji"]}',
        f'○ 내    용 : {data["content"]}',
    ]
    for i, txt in enumerate(items):
        add_tb(slide, OV_X, OV_Y + i*LINE_H, OV_W, LINE_H,
               text=txt, sz=11, bold=True)


    # ── 발생유형 + 결재 표 (두 개 분리) ──
    SEC1_Y = 1.59/2.54
    SEC1_H = (0.86 + 2.49 + (2.47 - 1.59 - 0.86)) / 2.54
    TX = OV_X + OV_W + 0.05/2.54
    TOTAL_W = (33.5/2.54) - TX - 0.15/2.54
    ROW_H = SEC1_H / 3

    # [표1] 발생유형 (3행 × 2열)
    T1_W = TOTAL_W * 0.38
    t1 = slide.shapes.add_table(3, 2, px(TX), px(SEC1_Y), px(T1_W), px(SEC1_H)).table
    remove_table_style(t1)
    t1.columns[0].width = px(T1_W * 0.55)
    t1.columns[1].width = px(T1_W * 0.45)
    for i in range(3): t1.rows[i].height = px(ROW_H)

    # 헤더: 발생유형 (2열 병합)
    t1.cell(0,0).merge(t1.cell(0,1))
    cell_write(t1.cell(0,0), '발생유형', sz=8, bold=True, align=PP_ALIGN.CENTER)
    set_bg(t1.cell(0,0), GRAY_HDR); set_border(t1.cell(0,0))

    fm = '◎' if data.get('type_first') else ''
    cell_write(t1.cell(1,0), '최초', sz=8, align=PP_ALIGN.CENTER)
    set_bg(t1.cell(1,0), WHITE); set_border(t1.cell(1,0))
    cell_write(t1.cell(1,1), fm, sz=8, align=PP_ALIGN.CENTER)
    set_bg(t1.cell(1,1), WHITE); set_border(t1.cell(1,1))

    rm = '◎' if data.get('type_recur') else ''
    cell_write(t1.cell(2,0), '재발', sz=8, align=PP_ALIGN.CENTER)
    set_bg(t1.cell(2,0), WHITE); set_border(t1.cell(2,0))
    cell_write(t1.cell(2,1), rm, sz=8, align=PP_ALIGN.CENTER)
    set_bg(t1.cell(2,1), WHITE); set_border(t1.cell(2,1))

    # [표2] 결재 (2행 × 4열)
    T2_X = TX + T1_W
    T2_W = TOTAL_W - T1_W
    t2 = slide.shapes.add_table(2, 4, px(T2_X), px(SEC1_Y), px(T2_W), px(SEC1_H)).table
    remove_table_style(t2)
    cw = T2_W / 4
    for i in range(4): t2.columns[i].width = px(cw)
    t2.rows[0].height = px(ROW_H)
    t2.rows[1].height = px(SEC1_H - ROW_H)

    # 헤더행: (빈) | P/L | 팀장 | 공장장
    cell_write(t2.cell(0,0), '', sz=8)
    set_bg(t2.cell(0,0), WHITE); set_border(t2.cell(0,0))
    for c, txt in enumerate(['P/L','팀장','공장장'], 1):
        cell_write(t2.cell(0,c), txt, sz=8, bold=True, align=PP_ALIGN.CENTER)
        set_bg(t2.cell(0,c), GRAY_HDR); set_border(t2.cell(0,c))

    # 결재행
    cell_write(t2.cell(1,0), '결재', sz=8, bold=True, align=PP_ALIGN.CENTER)
    set_bg(t2.cell(1,0), GRAY_HDR); set_border(t2.cell(1,0))
    for c, key in enumerate(['approval_pl','approval_team','approval_factory'], 1):
        cell_write(t2.cell(1,c), data.get(key,''), sz=8, align=PP_ALIGN.CENTER)
        set_bg(t2.cell(1,c), WHITE); set_border(t2.cell(1,c))

    # ────────────────────────────────────────────────────
    # 3. ■ 세부내용
    # ────────────────────────────────────────────────────
    SEC2_Y = SEC1_Y + SEC1_H + 0.07

    LBL2_W = 6.06/2.54
    add_rect(slide, M, SEC2_Y, LBL2_W, 0.86/2.54,
             fill=WHITE, line=None)
    add_tb(slide, M+0.05, SEC2_Y+0.03, LBL2_W-0.08, 0.72/2.54,
           text='■ 세부내용', sz=14, bold=True)

    # 세부내용 표
    DY = SEC2_Y + 0.3
    DH = SH - DY - M
    DW = SW - 2*M

    GU_W   = 0.65
    CONT_W = (DW - GU_W) / 2

    # 행: 헤더(0) / 현황(1) / 기타(2) / 향후(3)
    HEADER_ROW_H = 0.27
    GITA_ROW_H   = 0.30
    FUTURE_ROW_H = 0.38
    HYUNHWANG_H  = DH - HEADER_ROW_H - GITA_ROW_H - FUTURE_ROW_H

    dt = slide.shapes.add_table(4, 3, px(M), px(DY), px(DW), px(DH)).table
    remove_table_style(dt)

    dt.columns[0].width = px(GU_W)
    dt.columns[1].width = px(CONT_W)
    dt.columns[2].width = px(CONT_W)

    dt.rows[0].height = px(HEADER_ROW_H)
    dt.rows[1].height = px(HYUNHWANG_H)
    dt.rows[2].height = px(GITA_ROW_H)
    dt.rows[3].height = px(FUTURE_ROW_H)

    # 헤더 행
    for c, txt in enumerate(['구 분', '발 생 현 황', '발 생 원 인']):
        cell_write(dt.cell(0,c), txt, sz=9, bold=True, align=PP_ALIGN.CENTER)
        set_bg(dt.cell(0,c), GRAY_HDR); set_border(dt.cell(0,c))

    # ── 현황 행 ──
    cell_write(dt.cell(1,0), '현황', sz=9, bold=True, align=PP_ALIGN.CENTER)
    set_bg(dt.cell(1,0), GRAY_HDR); set_border(dt.cell(1,0))
    dt.cell(1,0).text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    from pptx.enum.text import MSO_ANCHOR
    dt.cell(1,0)._tc.get_or_add_tcPr()
    dt.cell(1,0).vertical_anchor = MSO_ANCHOR.MIDDLE

    # 발생 현황 (왼쪽): 현상 + 점검및조치
    sit_lines = [('■ 현상', 9, True)]
    for ln in data['situation'].split('\n'):
        if ln.strip():
            sit_lines.append((f'-{ln.strip()}', 9, False))
    if data.get('action','').strip():
        sit_lines.append(('', 5, False))
        sit_lines.append(('■ 점검 및 조치', 9, True))
        for ln in data['action'].split('\n'):
            if ln.strip():
                sit_lines.append((f'-{ln.strip()}', 9, False))
    cell_write(dt.cell(1,1), sit_lines, color=BLACK)
    set_bg(dt.cell(1,1), WHITE); set_border(dt.cell(1,1))

    # 발생 원인 (오른쪽): 텍스트 + 사진
    cause_lines = [('■ 발생 원인', 9, True)]
    for ln in data['cause'].split('\n'):
        if ln.strip():
            cause_lines.append((f'▶{ln.strip()}', 8.5, False))
    cell_write(dt.cell(1,2), cause_lines, color=BLACK)
    set_bg(dt.cell(1,2), WHITE); set_border(dt.cell(1,2))

    # ── 기타 행 ──
    cell_write(dt.cell(2,0), '기타', sz=9, bold=True, align=PP_ALIGN.CENTER)
    set_bg(dt.cell(2,0), GRAY_HDR); set_border(dt.cell(2,0))
    cell_write(dt.cell(2,1), f'제품처리 : {data.get("product_handling","")}', sz=8.5)
    set_bg(dt.cell(2,1), WHITE); set_border(dt.cell(2,1))
    cell_write(dt.cell(2,2), f'손실비용 : {data.get("loss_cost","")}', sz=8.5)
    set_bg(dt.cell(2,2), WHITE); set_border(dt.cell(2,2))

    # ── 향후 대응방안 행 ──
    cell_write(dt.cell(3,0), '향후\n대응방안', sz=9, bold=True,
               align=PP_ALIGN.CENTER)
    set_bg(dt.cell(3,0), GRAY_HDR); set_border(dt.cell(3,0))

    dt.cell(3,1).merge(dt.cell(3,2))
    future_lines = []
    for i, ln in enumerate(data['future'].split('\n')):
        if ln.strip():
            future_lines.append((f'{i+1}. {ln.strip()}', 8.5, False))
    cell_write(dt.cell(3,1), future_lines)
    set_bg(dt.cell(3,1), WHITE); set_border(dt.cell(3,1))

    # ────────────────────────────────────────────────────
    # 4. 사진 삽입 (표 위에 float)
    # ────────────────────────────────────────────────────
    # 현황 행 위치
    cell_top = DY + HEADER_ROW_H
    cell_h   = HYUNHWANG_H

    # 텍스트 줄수 × 줄높이(인치) 로 텍스트 영역 추정
    def txt_height(lines, line_h=0.175):
        return len(lines) * line_h + 0.1

    # 현장 사진 → 왼쪽(발생현황) 셀 하단
    if image_paths:
        th = txt_height(sit_lines)
        i_top = cell_top + th
        i_h   = max(cell_h - th - 0.05, 0.4)
        n = min(len(image_paths), 2)
        each_w = (CONT_W - 0.12) / n
        for idx, p in enumerate(image_paths[:2]):
            insert_pic(slide, p,
                       M + GU_W + 0.04 + idx*(each_w+0.06),
                       i_top, each_w, i_h)

    # 원인 사진 → 오른쪽(발생원인) 셀 하단
    if cause_image_paths:
        th = txt_height(cause_lines)
        i_top = cell_top + th
        i_h   = max(cell_h - th - 0.05, 0.4)
        n = min(len(cause_image_paths), 2)
        each_w = (CONT_W - 0.12) / n
        for idx, p in enumerate(cause_image_paths[:2]):
            insert_pic(slide, p,
                       M + GU_W + CONT_W + 0.04 + idx*(each_w+0.06),
                       i_top, each_w, i_h)

    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)
    return buf.read()


def cell_like_write(tf, text, sz=9, bold=False, color=BLACK,
                    align=PP_ALIGN.LEFT, font='맑은 고딕'):
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text; r.font.size = Pt(sz)
    r.font.bold = bold; r.font.color.rgb = color; r.font.name = font


# ── Flask ────────────────────────────────────────────────
@app.route('/')
def index():
    return render_template('index.html')


@app.route('/generate', methods=['POST'])
def generate():
    def save_files(file_list):
        paths = []
        for f in file_list:
            if f and f.filename:
                ext = os.path.splitext(f.filename)[1] or '.jpg'
                tmp = tempfile.NamedTemporaryFile(delete=False, suffix=ext,
                                                  dir=UPLOAD_FOLDER)
                f.save(tmp.name); paths.append(tmp.name)
        return paths

    image_paths       = save_files(request.files.getlist('images'))
    cause_image_paths = save_files(request.files.getlist('cause_images'))
    form = request.form

    data = {
        'title':      form.get('title',''),
        'team':       form.get('team',''),
        'date_str':   form.get('date_str',''),
        'datetime':   form.get('datetime',''),
        'jijong':     form.get('jijong',''),
        'jungji':     form.get('jungji',''),
        'content':    form.get('content',''),
        'reporter':   form.get('reporter',''),
        'type_first':        form.get('type_first') == 'on',
        'type_recur':        form.get('type_recur') == 'on',
        'approval_pl':       form.get('approval_pl',''),
        'approval_team':     form.get('approval_team',''),
        'approval_factory':  form.get('approval_factory',''),
        'situation':         form.get('situation',''),
        'cause':             form.get('cause',''),
        'action':            form.get('action',''),
        'product_handling':  form.get('product_handling',''),
        'loss_cost':         form.get('loss_cost',''),
        'future':            form.get('future',''),
    }

    ppt_bytes = create_report_ppt(data, image_paths, cause_image_paths)

    email_to = form.get('email_to','').strip()
    if email_to:
        try: send_email(email_to, ppt_bytes, data['title'])
        except Exception as e: print(f'메일 오류: {e}')

    for p in image_paths + cause_image_paths:
        try: os.unlink(p)
        except: pass

    fname = f"공정사고보고서_{data['date_str'].replace('.','').replace(' ','')}.pptx"
    return send_file(io.BytesIO(ppt_bytes), as_attachment=True,
                     download_name=fname,
                     mimetype='application/vnd.openxmlformats-officedocument'
                              '.presentationml.presentation')


def send_email(to_addr, ppt_bytes, title):
    import email.message
    from_addr = os.environ.get('MAIL_USER','')
    password  = os.environ.get('MAIL_PASS','')
    if not from_addr: raise ValueError('MAIL_USER 환경변수를 설정해주세요.')
    msg = email.message.EmailMessage()
    msg['From'] = from_addr
    msg['To'] = to_addr
    msg['Subject'] = f'[공정사고보고서] {title}'
    msg.set_content('공정사고 보고서를 첨부합니다.')
    msg.add_attachment(ppt_bytes,
                       maintype='application',
                       subtype='octet-stream',
                       filename='report.pptx')
    with smtplib.SMTP_SSL('smtp.gmail.com', 465) as s:
        s.login(from_addr, password)
        s.send_message(msg)


if __name__ == '__main__':
    app.run(debug=True, host='0.0.0.0', port=5000)
